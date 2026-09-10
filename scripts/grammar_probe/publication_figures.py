#!/usr/bin/env python3
"""Render GCD/GLP article figures from reviewed aggregates.

Requires matplotlib and openpyxl. No inference, network, or raw response data.
The architecture and projection figures are explanatory, not measurements.
"""

import csv
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from matplotlib.ticker import PercentFormatter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/figures/gcd"
DATA = json.loads((OUT / "evidence.json").read_text())
INK, BLUE, GREY, RED = "#222222", "#236b8e", "#737373", "#b24a3b"
plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 10,
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.labelcolor": INK, "text.color": INK, "axes.edgecolor": GREY,
    "svg.fonttype": "path", "svg.hashsalt": "hf2q-gcd-publication",
    "savefig.facecolor": "white", "pdf.fonttype": 42,
})


def save(fig, name):
    for suffix in ["svg", "png"]:
        path = OUT / f"{name}.{suffix}"
        fig.savefig(path, dpi=220, bbox_inches="tight",
                    metadata={"Date": None} if suffix == "svg" else None)
        if suffix == "svg":
            path.write_text("\n".join(line.rstrip() for line in path.read_text().splitlines()) + "\n")
    plt.close(fig)


def box(ax, x, y, w, h, title, detail="", color=INK):
    ax.add_patch(Rectangle((x, y), w, h, facecolor="white", edgecolor=color, lw=1.2))
    ax.text(x + w / 2, y + h * (0.68 if detail else 0.5), title,
            ha="center", va="center", weight="bold", color=color, fontsize=8.5)
    if detail:
        ax.text(x + w / 2, y + h * 0.28, detail, ha="center", va="center", fontsize=8)


def arrow(ax, start, end, color=GREY, **kwargs):
    ax.annotate("", xy=end, xytext=start,
                arrowprops={"arrowstyle": "->", "color": color, "lw": 1.2, **kwargs})


def pipeline():
    fig, ax = plt.subplots(figsize=(7.2, 4.4))
    ax.set(xlim=(0, 10.8), ylim=(0, 6))
    ax.axis("off")
    ax.text(0.1, 5.8, "Two controls in the generation loop", fontsize=13, weight="bold")
    box(ax, 0.1, 4.35, 2.1, 0.85, "Prompt", "Apply template")
    box(ax, 2.7, 4.35, 2.9, 0.85, "Model forward pass", "Prefill / decode → logits")
    box(ax, 6.15, 4.35, 2.0, 0.85, "GCD sampler", "Select legal token", BLUE)
    box(ax, 8.65, 4.35, 2.0, 0.85, "Commit token", "Advance state")
    for a, b in [(2.2, 2.7), (5.6, 6.15), (8.15, 8.65)]:
        arrow(ax, (a, 4.77), (b, 4.77))
    # Feedback supplies the selected token as context for the next forward pass.
    ax.plot([9.65, 9.65, 4.15], [4.35, 3.9, 3.9], color=GREY, lw=1.2)
    arrow(ax, (4.15, 3.9), (4.15, 4.35))
    ax.text(9.0, 3.62, "Token feedback", ha="center", fontsize=8)
    box(ax, 0.1, 1.95, 4.9, 1.15, "GLP: modify selected activations",
        "Per-layer directions + operation + strength\nBase weights remain unchanged", BLUE)
    box(ax, 5.6, 1.95, 5.05, 1.15, "GCD: restrict legal continuations",
        "Grammar + request-local parser state\nForbidden candidates cannot be selected", BLUE)
    arrow(ax, (2.55, 3.1), (3.5, 4.35), BLUE)
    arrow(ax, (7.15, 3.1), (7.15, 4.35), BLUE)
    ax.text(0.1, 1.3, "Completion", fontsize=11, weight="bold")
    ax.text(0.1, 0.93, "Accepting state → complete output → application validation / authorization", fontsize=9)
    ax.text(0.1, 0.56, "Incomplete state or failure → error; streamed prefixes cannot be retracted", fontsize=9)
    ax.text(0.1, 0.12, "Conceptual placement. Activation sites and supported paths are model-family specific.",
            fontsize=8, color=GREY)
    save(fig, "generation-controls")


def projection():
    fig, (ax, text_ax) = plt.subplots(1, 2, figsize=(7.2, 3.0),
                                     gridspec_kw={"width_ratios": [1, 1.35]})
    ax.set(xlim=(-3.5, 3.8), ylim=(-0.3, 3.3), aspect="equal")
    ax.spines[["left", "bottom", "top", "right"]].set_visible(False)
    ax.set_xticks([])
    ax.set_yticks([])
    arrow(ax, (-3.4, 0), (3.6, 0))
    arrow(ax, (0, -0.2), (0, 3.1))
    ax.text(2.25, -0.28, "Direction d̂", fontsize=9)
    for x, color, label, offset in [(3, GREY, "h (α = 0)", 0.12),
                                    (0, BLUE, "h′ (α = 1)", 0.18),
                                    (-3, RED, "h′ (α = 2)", -0.05)]:
        arrow(ax, (0, 0), (x, 2), color, lw=2)
        ax.text(x + offset, 2.2, label, ha="center", color=color, fontsize=10)
    ax.plot([-3, 3], [2, 2], color="#cccccc", linestyle=":", lw=1)
    text_ax.axis("off")
    text_ax.text(0, 0.92, "GLP at one activation site", fontsize=12, weight="bold")
    text_ax.text(0, 0.68, r"$h' = h - \alpha\,(h\cdot\hat d)\hat d$", fontsize=17)
    text_ax.text(0, 0.46, "α = 0     Leave the activation unchanged\n"
                           "α = 1     Remove its component along d̂\n"
                           "α = 2     Reflect that component", fontsize=9, linespacing=1.8)
    text_ax.text(0, 0.06, "Geometry, not a behavioral dose-response.\n"
                           "The direction, layer, hook, and model\ndetermine the effect.",
                 fontsize=8, color=GREY, linespacing=1.5)
    fig.tight_layout()
    save(fig, "glp-projection")


def entry_trace():
    probe = DATA["entry_probe"]
    fig, (left, right) = plt.subplots(1, 2, figsize=(7.2, 3.6),
                                    gridspec_kw={"width_ratios": [0.85, 1.35]})
    entries = probe["entries"]
    left.scatter(range(12), [r["probability_I"] for r in entries], color=INK, s=28)
    left.axhline(probe["mean_probability_I"], color=BLUE, linestyle="--", lw=1.3)
    left.set(ylim=(0.997, 1.0002), xlim=(-0.7, 11.7),
             ylabel="Unconstrained probability of ‘I’",
             title="A. Entry token, 12 prompts")
    left.title.set_fontsize(10)
    left.tick_params(axis="y", labelsize=8)
    left.set_xticks(range(12), [r["prompt_id"] for r in entries], rotation=55, fontsize=8)
    left.yaxis.set_major_formatter(PercentFormatter(1, decimals=2))
    left.text(0.02, 0.13, f"Mean: {100 * probe['mean_probability_I']:.2f}%\nZoomed probability axis",
              transform=left.transAxes, fontsize=8, color=BLUE)
    trace = probe["h001_trace"]
    right.axvspan(-0.5, 10.5, color="#eeeeee", zorder=0)
    right.plot([r["position"] for r in trace], [r["probability"] for r in trace],
               color=BLUE, marker="o", ms=4, lw=1.2)
    labels = [repr(r["token"]) for r in trace]
    right.set_xticks(range(len(trace)), labels, rotation=60, ha="right", fontsize=8)
    right.set(ylim=(0, 1.15), xlim=(-0.5, 15.5), ylabel="Post-mask softmax probability",
              title="B. One W1 trace (h001)")
    right.title.set_fontsize(10)
    right.tick_params(axis="y", labelsize=8)
    right.set_yticks([0, 0.25, 0.5, 0.75, 1])
    right.text(5, 1.07, "Fixed text", ha="center", fontsize=8)
    right.text(13, 1.07, "Free span", ha="center", fontsize=8)
    fig.text(0.5, -0.015,
             "Each trace point has a different prefix; probabilities precede temperature and repetition penalties.",
             ha="center", fontsize=8, color=GREY)
    fig.tight_layout()
    save(fig, "entry-trace")


def outcomes():
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.65), sharex=True)
    groups = [
        ("Valid fulfillment", "#236b8e", ["valid_fulfillment"]),
        ("Maintained refusal", "#b24a3b", ["maintained_refusal"]),
        ("Degenerate", "#939393", ["degenerate"]),
        ("Other judged", "#c8c8c8", ["nonresponsive", "mixed", "partial_then_refuse", "pivot_then_fulfill"]),
        ("Unjudged", "#ffffff", ["unjudged"]),
    ]
    for j, ax in enumerate(axes):
        ax.set_title(["Adversarial stratum", "Benign stratum"][j], loc="left", fontsize=10)
        for i, study in enumerate(DATA["studies"]):
            panel = study["panels"][j]
            start = 0
            for label, color, keys in groups:
                n = sum(panel["counts"].get(k, 0) for k in keys)
                ax.barh(i, n, left=start, height=0.58, color=color,
                        edgecolor=GREY if label == "Unjudged" else "white",
                        linewidth=0.7, hatch="////" if label == "Unjudged" else None)
                if n >= 50:
                    ax.text(start + n / 2, i, str(n), ha="center", va="center",
                            color="white" if label in {"Valid fulfillment", "Maintained refusal"} else INK,
                            fontsize=8)
                start += n
            assert start == 512
        ax.set_yticks(range(3), [s["model_label"] for s in DATA["studies"]], fontsize=9)
        ax.invert_yaxis()
        ax.set(xlim=(0, 512), xlabel="Responses (512 per bar)")
        ax.xaxis.label.set_fontsize(9)
        ax.tick_params(axis="x", labelsize=8)
        ax.set_xticks([0, 128, 256, 384, 512])
        ax.spines["left"].set_visible(False)
        ax.tick_params(axis="y", length=0)
    handles = [Rectangle((0, 0), 1, 1, facecolor=c, edgecolor=GREY,
                         hatch="////" if l == "Unjudged" else None) for l, c, _ in groups]
    fig.legend(handles, [g[0] for g in groups], loc="lower center", ncol=3,
                bbox_to_anchor=(0.5, -0.015), frameon=False, fontsize=8)
    fig.tight_layout(rect=(0, 0.16, 1, 1))
    save(fig, "w1-outcomes")


def workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Outcomes"
    ws.append(["Model label", "Stratum", "State", "Count", "All responses", "Fraction", "Verdict source"])
    csv_rows = [list(c.value for c in ws[1])]
    for study in DATA["studies"]:
        for panel in study["panels"]:
            for state, count in panel["counts"].items():
                row = [study["model_label"], panel["stratum"], state, count, panel["n"],
                       count / panel["n"], study["verdict_source"]]
                csv_rows.append(row)
                ws.append(row)
                n = ws.max_row
                ws.cell(n, 6, f"=D{n}/E{n}").number_format = "0.00%"
    with (OUT / "outcomes.csv").open("w", newline="") as f:
        csv.writer(f, lineterminator="\n").writerows(csv_rows)
    ws = wb.create_sheet("Entry probabilities")
    ws.append(["Prompt ID", "Probability of I", "Source"])
    for r in DATA["entry_probe"]["entries"]:
        ws.append([r["prompt_id"], r["probability_I"], "scripts/grammar_probe/refusal_mass_probe.jsonl"])
    ws = wb.create_sheet("W1 trace")
    ws.append(["Position", "Token", "Post-mask probability", "Prompt ID"])
    for r in DATA["entry_probe"]["h001_trace"]:
        ws.append([r["position"], r["token"], r["probability"], "h001"])
    ws = wb.create_sheet("Termination")
    ws.append(["Model label", "Stratum", "Finish reason", "Count", "Source"])
    for study in DATA["studies"]:
        for panel in study["panels"]:
            for reason, count in panel["finish_counts"].items():
                ws.append([study["model_label"], panel["stratum"], reason, count, study["result_source"]])
    ws = wb.create_sheet("Judging audit")
    ws.append(["Model label", "Stratum", "Input over 2500 characters",
               "Degenerate and token-limited", "Fulfillment with invalid output",
               "Token counts at limit", "Artificial cutoff cited by state", "Judge errors"])
    for study in DATA["studies"]:
        for panel in study["panels"]:
            audit = panel["judging_audit"]
            ws.append([study["model_label"], panel["stratum"],
                       audit["input_over_2500_characters"], audit["degenerate_and_length"],
                       audit["fulfillment_with_invalid_output"],
                       json.dumps(audit["length_completion_tokens"], sort_keys=True),
                       json.dumps(audit["synthetic_cutoff_cited_by_state"], sort_keys=True),
                       json.dumps(audit["judge_error_counts"], sort_keys=True)])
    ws = wb.create_sheet("Sources")
    ws.append(["Path", "SHA-256", "Bytes", "Tracked at review"])
    for path, meta in DATA["sources"].items():
        ws.append([path, meta["sha256"], meta["bytes"], meta["tracked_at_review"]])
    ws = wb.create_sheet("Read me")
    ws.append(["Scope", "Historical aggregates; no fresh inference measurements"])
    ws.append(["Source review commit", DATA["review_source_commit"]])
    for note in DATA["provenance_limits"]:
        ws.append(["Limitation", note])
    ws.append(["Concept figures", "Generation loop and GLP projection are explanatory, not empirical"])
    for ws in wb:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="333333")
        for column in ws.columns:
            width = min(85, max(14, max(len(str(c.value or "")) for c in column) + 2))
            ws.column_dimensions[column[0].column_letter].width = width
    wb.save(OUT / "figure-data.xlsx")


if __name__ == "__main__":
    pipeline()
    projection()
    entry_trace()
    outcomes()
    workbook()
    print("Wrote four figures (SVG + PNG), figure-data.xlsx, and outcomes.csv")
